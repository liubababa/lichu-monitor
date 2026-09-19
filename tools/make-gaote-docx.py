#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
生成《高特 CCU 接入参数》Word 文档（发给高特，对应手册 4.1 平台配置）。

用法：python tools/make-gaote-docx.py <平台主机地址> <端口> <账号> <密码>
     密码等敏感信息通过命令行传入，不写进代码，避免随仓库泄露。
输出：给高特的接入参数.docx（项目根目录）
"""
import os
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '给高特的接入参数.docx')
BODY = '宋体'
HEAD = '黑体'

BROKER_HOST = sys.argv[1] if len(sys.argv) > 1 else 'mqtt.ykdesign.top'
BROKER_PORT = sys.argv[2] if len(sys.argv) > 2 else '1883'
CCU_USER = sys.argv[3] if len(sys.argv) > 3 else 'gaote_ccu'
CCU_PASS = sys.argv[4] if len(sys.argv) > 4 else '<密码>'


def run(p, text, size=11, bold=False, font=BODY):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = RGBColor(0, 0, 0)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    return r


def para(doc, text='', size=11, bold=False, font=BODY, indent=True,
         align=None, space_after=6, space_before=0, line=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(22)
    if align is not None:
        p.alignment = align
    if text:
        run(p, text, size, bold, font)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.6)
    sec.top_margin = sec.bottom_margin = Cm(2.4)

    st = doc.styles['Normal']
    st.font.name = BODY
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(16)
    run(t, '高特 CCU 接入参数', 16, True, HEAD)

    para(doc, '您好，')
    para(doc, '我们这边已经按手册第 4.1 节“平台配置与连接”准备好了对接参数，'
              '麻烦在 CCU 的物联网平台配置界面按下面填写（端口是 1883 明文，不需要证书、不需要 TLS）：')

    rows = [
        ('物联网启用', '启用'),
        ('物联网平台名称', '力储未来监控平台（只作显示用，可随意填）'),
        ('物联网平台ID', '0（保持默认，不要修改）'),
        ('平台主机地址', BROKER_HOST),
        ('端口', BROKER_PORT + '（TCP 明文，不做 TLS）'),
        ('账号', CCU_USER),
        ('密码', CCU_PASS),
        ('用户ID', '保持默认 {deviceId}，我方无特殊要求'),
        ('密钥', '保持默认 RFyim2I3IcfK（手册说明不能随意修改）'),
    ]
    tb = doc.add_table(rows=0, cols=2)
    tb.style = 'Table Grid'
    for k, v in rows:
        c = tb.add_row().cells
        c[0].text = ''
        c[1].text = ''
        run(c[0].paragraphs[0], k, 10.5, True)
        run(c[1].paragraphs[0], v, 10.5)
        c[0].width = Cm(3.6)
        c[1].width = Cm(11.8)

    para(doc, '', space_after=4)
    para(doc, '连上以后我们会订阅 /{ProductSN}/{DeviceSN}/#，也就是说：')
    para(doc, '· 遥测、遥信：/{ProductSN}/{DeviceSN}/rtg/data|status/{维度}/…（30 秒周期）', indent=True, space_after=2)
    para(doc, '· 遥调、遥控回读：/{ProductSN}/{DeviceSN}/cmd/get/emu/{功能}（我们下发 cmd/set/emu/{功能} 后，'
              '请按手册 4.3 立即触发回读）', indent=True, space_after=2)
    para(doc, '· 断网补传：/{ProductSN}/{DeviceSN}/history/…（按手册 4.4）', indent=True, space_after=6)

    para(doc, '主题与点位按本项目的协议文档 SJ2025B3781ESCCU-MQTT 执行，不需要额外改动。'
              '另外遥调、遥控下发前需要现场把控制源切到远程模式（手册 4.3），到时候麻烦现场配合一下。')

    para(doc, '有四件事想请你们确认：', space_before=6)

    para(doc, '1、这台 CCU 的 ProductSN 和 DeviceSN 实际值是多少？Topic 里要用。'
              '手册 5.1 节里的 kp23bhcpmt91n2v8、S230612B0125 是示例项目的值，'
              '同一个手册里还出现过 R221010B2001，我们不确定现场是哪个，麻烦给准确值。', indent=True)
    para(doc, '2、现场固件实际发的主题是哪种形状？堆、簇是 array/cluster 全写还是手册示例的 arr/clu 简写？'
              '电表主题是否带 /extend/{厂商型号} 段（现场装了哪几块表、对应哪个型号）？'
              '我们两种都做了兼容，但想确认以哪个为准，最好能抓一条真实报文给我们。', indent=True)
    para(doc, '3、压缩有没有启用？我们目前按未压缩的 JSON 报文解析，'
              '如果主题里带 comp-gzip / comp-lz4，请提前说一声，我们这边补上解压处理。', indent=True)
    para(doc, '4、CCU 是否支持 TLS（8883）？现在先用 1883 明文对接，'
              '如果固件支持加密连接，我们这边可以再开加密端口，安全性更好。', indent=True)

    para(doc, '另外提醒一句：1883 是明文传输，账号密码和报文在公网是可见的，请不要外传；'
              '如果储能站那边的出口是固定公网 IP，告诉我们，我们可以在服务器上把 1883 '
              '限制成只允许这个 IP 访问，更保险。（这一项可选，不做也能正常连。）')

    para(doc, '我们这边已经用模拟设备把协议里定义的主题跑过一轮：遥测、遥信、下发与回读、'
              '历史补传都能收，页面显示和下发都正常。你们配置好之后说一声，我们在线盯一下报文，'
              '确认没问题就正式跑起来。')

    para(doc, '有不清楚的地方随时联系。')

    para(doc, '', space_after=10)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, '联系人：', 11)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, '电　话：', 11)

    doc.save(OUT)
    print('已生成：' + os.path.abspath(OUT))


if __name__ == '__main__':
    main()
