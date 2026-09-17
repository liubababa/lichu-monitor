#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
解析 MQTT 协议文档 xlsx（厂家给的点表/主题定义），输出可读文本。

用法：python tools/parse-protocol-xlsx.py <xlsx路径> [输出txt路径]
"""
import os
import sys

import openpyxl

if len(sys.argv) < 2:
    print('用法：python tools/parse-protocol-xlsx.py <xlsx路径> [输出txt路径]')
    sys.exit(1)

src = sys.argv[1]
dst = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + '.txt'

wb = openpyxl.load_workbook(src, data_only=True)
print('工作表：', ' | '.join(wb.sheetnames))

out = []
for ws in wb.worksheets:
    out.append('\n' + '=' * 70)
    out.append('工作表：%s   行数 %d  列数 %d' % (ws.title, ws.max_row, ws.max_column))
    out.append('=' * 70)
    for row in ws.iter_rows(values_only=True):
        cells = ['' if c is None else str(c).strip() for c in row]
        while cells and cells[-1] == '':
            cells.pop()
        if not cells:
            continue
        out.append(' | '.join(cells))

text = '\n'.join(out)
with open(dst, 'w', encoding='utf-8') as f:
    f.write(text)
print('已写入：%s（%d 字符）' % (os.path.abspath(dst), len(text)))
