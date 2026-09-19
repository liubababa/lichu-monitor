#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
生成《晶农EMS 接入参数》Word 文档（发给厂家）。

用法：python tools/make-vendor-docx.py <broker地址> <端口> <用户名> <密码>
     密码等敏感信息通过命令行传入，不写进代码，避免随仓库泄露。
输出：给厂家的接入参数.docx（项目根目录）
"""
import os
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '给厂家的接入参数.docx')
BODY = '宋体'
HEAD = '黑体'

BROKER_HOST = sys.argv[1] if len(sys.argv) > 1 else 'mqtt.ykdesign.top'
BROKER_PORT = sys.argv[2] if len(sys.argv) > 2 else '1883'
EMS_USER = sys.argv[3] if len(sys.argv) > 3 else 'zhhn_ems'
EMS_PASS = sys.argv[4] if len(sys.argv) > 4 else '<密码>'


def run(p, text, size=11, bold=False, font=BODY):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = RGBColor(0, 0, 0)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    return r


def para(doc, text='', size=11, bold=False, font=BODY, indent=True,
         align=None, space_after=6, space_before=0, line=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(22)
    if align is not None:
        p.alignment = align
    if text:
        run(p, text, size, bold, font)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.6)
    sec.top_margin = sec.bottom_margin = Cm(2.4)

    st = doc.styles['Normal']
    st.font.name = BODY
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

    # 标题
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(16)
    run(t, '晶农EMS 接入参数', 16, True, HEAD)

    para(doc, '您好，')
    para(doc, '按你们反馈的“EMS 只支持 1883 明文”，我们这边的 MQTT 服务器已经改成 1883 明文端口'
              '重新搭好了，不需要证书、不需要 TLS。麻烦在 EMS 上按下表配置第三方 broker：')

    # 参数表（朴素样式）
    rows = [
        ('服务器地址', BROKER_HOST),
        ('端口', BROKER_PORT + '（TCP 明文，不做 TLS）'),
        ('用户名', EMS_USER),
        ('密码', EMS_PASS),
        ('QoS', '0 或 1 都可以'),
        ('Keepalive', '60 秒'),
        ('ClientId', '填 EMS 自己的 SN 即可'),
        ('会话', 'Clean Session = true（不需要保留离线消息）'),
    ]
    tb = doc.add_table(rows=0, cols=2)
    tb.style = 'Table Grid'
    for k, v in rows:
        c = tb.add_row().cells
        c[0].text = ''
        c[1].text = ''
        run(c[0].paragraphs[0], k, 10.5, True)
        run(c[1].paragraphs[0], v, 10.5)
        c[0].width = Cm(3.4)
        c[1].width = Cm(12.0)

    para(doc, '', space_after=4)
    para(doc, '登录和上报的 Topic 按协议文档来，就是 zhhn/Post/Login/{SN} 和 '
              'zhhn/Post/PeriodReport/{SN}。我们收到登录请求会回 zhhn/PostRsp/Login/{SN}，result 置 1——'
              '这一步我们已经用模拟设备测通了，EMS 发登录后会正常收到应答。')

    para(doc, '地址这边我们填的是域名而不是 IP：万一以后我们换服务器，只要改解析，你们那边不用动配置。')

    para(doc, '还有两件事想跟你们确认：', space_before=6)

    para(doc, '1、这台 EMS 的 SN 是多少？Topic 里要用，我们好提前把设备台账对上。', indent=True)
    para(doc, '2、通道下发里的 wayName 该怎么填？文档示例里是“下设充电/放电功率”这种中文名，'
              '点表里又是 SetPower、S1Power 这样的英文 TAG，两边对不上。麻烦给一份对照表'
              '（包括 S1~S10 分段策略那几个字段），不然充放电策略我们下发不了。', indent=True)

    para(doc, '另外提醒一句：这端口是明文传输，账号密码和报文在公网是可见的，'
              '所以现场别配错账号、也别把密码外传；如果储能站那边的出口是固定公网 IP，'
              '告诉我们，我们可以在服务器上把 1883 限制成只允许这个 IP 访问，更保险。'
              '（这一项可选，不做也能正常连。）', space_before=6)

    para(doc, '还有，协议文档里有几处示例报文的写法有点问题，主要是全角冒号和逗号、'
              'PeriodReport 里的 time 写在了 data 数组里面、UserInfor 应答里 DeviceName 重复。'
              '我们平台做了兼容，按文档原样发过来也能收，但建议你们顺手改一下，'
              '具体哪几处我整理在附件的问题清单里了。')

    para(doc, '我们这边已经先测过一轮：文档里定义的 14 个 Topic 都能通，页面收数据、'
              '查设备清单、下发通道都正常。你们配置好之后跟我说一下，我们在线盯一下报文，'
              '确认没问题就正式跑起来。')

    para(doc, '有不清楚的地方随时联系。')

    para(doc, '', space_after=10)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, '联系人：', 11)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, '电　话：', 11)

    doc.save(OUT)
    print('已生成：' + os.path.abspath(OUT))


if __name__ == '__main__':
    main()
