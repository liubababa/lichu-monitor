#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
把高特 MQTT 协议文档 xlsx 生成为前端可用的点表 js（js/gaote-points.js）。

用法：python tools/gen-gaote-points.py refs/SJ2025B3781ESCCU-MQTT.xlsx
"""
import json
import os
import re
import sys

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else 'refs/SJ2025B3781ESCCU-MQTT.xlsx'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'js', 'gaote-points.js')


def txt(v):
    if v is None:
        return ''
    s = str(v).strip()
    return '' if s == 'None' else s


def find_header(ws):
    """找到含 'JSON/protobuf字段' 的表头行，返回 (行号, 列名->列索引)"""
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        cells = [txt(c) for c in row]
        if any('JSON/protobuf' in c for c in cells):
            idx = {}
            for j, c in enumerate(cells):
                if c:
                    idx[c] = j
            return i, idx
    return 0, {}


def pick(idx, *names):
    for n in names:
        for k, v in idx.items():
            if n in k:
                return v
    return None


def parse_enum(s):
    """从 '0-正常,1-故障' 或 '{"enum":{"0":"正常"}}' 里提取枚举"""
    if not s:
        return None
    m = re.search(r'"enum"\s*:\s*\{(.*?)\}', s.replace('\n', ' '))
    if m:
        pairs = re.findall(r'"([^"]+)"\s*:\s*"([^"]*)"', m.group(1))
        if pairs:
            return {k: v for k, v in pairs}
    pairs = re.findall(r'(\d+)\s*[-:：]\s*([^,;；]+)', s)
    if pairs and len(pairs) >= 2:
        return {k: v.strip() for k, v in pairs}
    return None


def dump(ws, kind):
    """解析遥测/遥信表：按数据维度分组，返回 {dim: [点位]}"""
    hdr_row, idx = find_header(ws)
    c_dim = pick(idx, '数据维度')
    c_subj = pick(idx, '设备主题')
    c_field = pick(idx, 'JSON/protobuf字段')
    c_key = pick(idx, '属性标识')
    c_name = pick(idx, '名称')
    c_type = pick(idx, '类型')
    c_unit = pick(idx, '单位')
    c_def = pick(idx, '数值功能定义')
    c_note = pick(idx, '备注')
    c_min = pick(idx, '最小值')
    c_max = pick(idx, '最大值')
    rows = list(ws.iter_rows(values_only=True))[hdr_row:]

    groups = {}
    dim = subj = ''
    for r in rows:
        if not r:
            continue
        get = lambda i: txt(r[i]) if i is not None and i < len(r) else ''
        if get(c_subj):
            subj = get(c_subj)
        if get(c_dim):
            dim = get(c_dim)
        field = get(c_field)
        name = get(c_name)
        if not field or not dim:
            continue
        try:
            i_field = int(float(field))
        except ValueError:
            continue
        item = {'i': i_field, 'n': name or get(c_key), 'k': get(c_key)}
        if get(c_type):
            item['t'] = get(c_type)
        if get(c_unit):
            item['u'] = get(c_unit)
        if get(c_min):
            item['min'] = get(c_min)
        if get(c_max):
            item['max'] = get(c_max)
        en = parse_enum(get(c_def) or get(c_note))
        if en:
            item['e'] = en
        if subj:
            item['src'] = subj
        groups.setdefault(dim, []).append(item)
    for d in groups:
        groups[d].sort(key=lambda x: x['i'])
    return groups


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)

    # 文档信息
    meta = {}
    for row in wb['更新'].iter_rows(values_only=True):
        if row and txt(row[0]):
            meta[txt(row[0])] = txt(row[1]) if len(row) > 1 else ''
    for row in wb['接入站点'].iter_rows(values_only=True):
        if row and txt(row[0]):
            meta[txt(row[0])] = txt(row[1]) if len(row) > 1 else ''

    # 主题定义
    topics = []
    ws = wb['主题定义']
    rows = list(ws.iter_rows(values_only=True))
    for r in rows[1:]:
        if not r or not txt(r[2]):
            continue
        topics.append({
            'ver': txt(r[0]), 'cat': txt(r[1]), 'desc': txt(r[2]), 'dim': txt(r[3]),
            'up': txt(r[4]), 'down': txt(r[5]), 'hist': txt(r[6]),
            'period': txt(r[9]), 'dev': txt(r[10]), 'cls': txt(r[12])
        })

    # 遥测 / 遥信 点位
    telem = dump(wb['遥测'], 'rtg')
    status = dump(wb['遥信'], 'sts')

    # 遥调 / 遥控（按主题分组）
    setpoints = {}
    for sheet in ['遥调', '遥控']:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        hdr_row, idx = find_header(ws)
        c_type = pick(idx, '功能类型')
        c_topic = pick(idx, '主题')
        c_no = pick(idx, '点号')
        c_name = pick(idx, '名称')
        c_dtype = pick(idx, '类型')
        c_json = pick(idx, 'JSON/protobuf字段')
        c_def = pick(idx, '数值功能定义')
        c_unit = pick(idx, '单位')
        rows = list(ws.iter_rows(values_only=True))[hdr_row:]
        ftype = topic = ''
        for r in rows:
            if not r:
                continue
            get = lambda i: txt(r[i]) if i is not None and i < len(r) else ''
            if get(c_type):
                ftype = get(c_type)
            if get(c_topic):
                topic = get(c_topic)
            if not topic or not get(c_no):
                continue
            try:
                no = int(float(get(c_no)))
            except ValueError:
                continue
            item = {'i': no, 'n': get(c_name), 'k': get(c_json), 't': get(c_dtype),
                    'sheet': sheet, 'func': ftype}
            if get(c_unit):
                item['u'] = get(c_unit)
            en = parse_enum(get(c_def))
            if en:
                item['e'] = en
            setpoints.setdefault(topic, []).append(item)

    # 设备拓扑
    topology = []
    ws = wb['设备拓扑']
    rows = list(ws.iter_rows(values_only=True))
    for r in rows[1:]:
        if not r or not txt(r[1]):
            continue
        topology.append({
            'dev': txt(r[1]), 'name': txt(r[2]), 'sub': txt(r[3]), 'subName': txt(r[4]),
            'arr': txt(r[5]), 'clu': txt(r[6]), 'idx': txt(r[7]),
            'mfr': txt(r[9]), 'model': txt(r[10]), 'ver': txt(r[17])
        })

    stat = {
        'telemDims': len(telem), 'telemPoints': sum(len(v) for v in telem.values()),
        'statusDims': len(status), 'statusPoints': sum(len(v) for v in status.values()),
        'setTopics': len(setpoints), 'setPoints': sum(len(v) for v in setpoints.values()),
        'topics': len(topics), 'devices': len(topology)
    }

    js = []
    js.append('/* 自动生成自 %s —— 请勿手工编辑，重新生成：python tools/gen-gaote-points.py */' % os.path.basename(SRC))
    js.append('window.GAOTE = {')
    js.append('  meta: %s,' % json.dumps(meta, ensure_ascii=False))
    js.append('  stat: %s,' % json.dumps(stat, ensure_ascii=False))
    js.append('  topics: %s,' % json.dumps(topics, ensure_ascii=False))
    js.append('  points: %s,' % json.dumps(telem, ensure_ascii=False))
    js.append('  status: %s,' % json.dumps(status, ensure_ascii=False))
    js.append('  setpoints: %s,' % json.dumps(setpoints, ensure_ascii=False))
    js.append('  topology: %s,' % json.dumps(topology, ensure_ascii=False))
    js.append('  /* 按数据维度 + 索引快速查点位 */')
    js.append('  get: function (dim, index, isStatus) {')
    js.append('    var g = (isStatus ? this.status : this.points)[dim];')
    js.append('    if (!g) return null;')
    js.append('    if (!g.__map) { var m = {}; for (var i = 0; i < g.length; i++) m[g[i].i] = g[i]; g.__map = m; }')
    js.append('    return g.__map[index] || null;')
    js.append('  }')
    js.append('};')

    with open(OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(js) + '\n')

    print('已生成：%s' % os.path.abspath(OUT))
    for k, v in stat.items():
        print('  %s: %s' % (k, v))
    print('  遥测维度：%s' % ', '.join(sorted(telem.keys())))
    print('  遥信维度：%s' % ', '.join(sorted(status.keys())))
    print('  遥调主题：%s' % ', '.join(sorted(setpoints.keys())))


if __name__ == '__main__':
    main()
