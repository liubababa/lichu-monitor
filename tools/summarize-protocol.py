#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""统计 MQTT 协议文档各工作表的点位数量与设备分布，输出摘要。"""
import sys
from collections import Counter, OrderedDict

import openpyxl

src = sys.argv[1] if len(sys.argv) > 1 else 'refs/SJ2025B3781ESCCU-MQTT.xlsx'
wb = openpyxl.load_workbook(src, data_only=True)

print('== 文档信息 ==')
ws = wb['更新']
for row in ws.iter_rows(values_only=True):
    if row and row[0]:
        print('  %s: %s' % (row[0], row[1] if len(row) > 1 else ''))
ws = wb['接入站点']
for row in ws.iter_rows(values_only=True):
    if row and row[0]:
        print('  %s: %s' % (row[0], row[1] if len(row) > 1 else ''))

print('\n== 主题定义（按类别）==')
ws = wb['主题定义']
rows = list(ws.iter_rows(values_only=True))
header = rows[0]
cat_i, desc_i, up_i, down_i = 1, 2, 4, 5
stat = OrderedDict()
for r in rows[1:]:
    if not r or not r[cat_i]:
        continue
    cat = str(r[cat_i])
    stat.setdefault(cat, []).append(str(r[desc_i] or ''))
for cat, items in stat.items():
    print('  %s：%d 条' % (cat, len(items)))
    for d in items[:6]:
        print('     · %s' % d)
    if len(items) > 6:
        print('     … 其余 %d 条' % (len(items) - 6))

for name in ['遥测', '遥信', '遥调', '遥控']:
    if name not in wb.sheetnames:
        continue
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    print('\n== %s 工作表（%d 行）==' % (name, len(rows) - 1))
    print('  列：', ' | '.join([str(h) for h in header if h]))
    # 设备主题列(1) 与 数据维度列(2) 统计（合并单元格为空则沿用上一个）
    dev = None
    dim = None
    cnt = Counter()
    for r in rows[1:]:
        if not r:
            continue
        if r[1]:
            dev = str(r[1]).strip()
        if len(r) > 2 and r[2]:
            dim = str(r[2]).strip()
        if len(r) > 3 and r[3] is not None:
            cnt[(dev or '?', dim or '?')] += 1
    for (d, m), n in cnt.most_common():
        print('  %-14s %-14s %3d 个点位' % (d, m, n))
